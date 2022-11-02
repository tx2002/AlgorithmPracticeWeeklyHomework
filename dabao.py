import sys
import time
import resource

from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.cell import _writer
from docx import Document
from faker import Faker
from PyQt5.QtCore import *
from PyQt5.QtGui import *
from PyQt5.QtWidgets import *

# pyinstaller -F -w dabao.py -p resource.py -i njupt.ico
def int_function(count, min, max, must_contain, must_not_contain):
    res = set()
    count_old = count
    # 判断约束是否正确
    if count < 0:
        return "生成个数不能为负数！"
    if min > max:
        return "最小值不能大于最大值！"

    contain_num = []
    if must_contain != "":
        contain_str = must_contain.split(',')
        contain_num = [int(i) for i in contain_str]
        for num in contain_num:
            if num > max or num < min:
                return "必须包含的数字不在约束范围内！"
        # 添加必须有的数字
        for num in contain_num:
            if num not in res:
                res.add(num)
                count = count - 1

    count_not_contain = 0
    if must_not_contain != "":
        not_contain_str = must_not_contain.split(',')
        not_contain_num = set()
        for i in not_contain_str:
            not_contain_num.add(int(i))
        for num in iter(not_contain_num):
            if num in contain_num:
                return "必须包含的数字与必须不包含的数字冲突！"
            if num >= min and num <= max:
                count_not_contain = count_not_contain + 1

    if max - min + 1 - count_not_contain < count_old:
        return "范围内没有" + str(count_old) + "个非重复数据！"

    faker = Faker(locale="zh_CN")

    while len(res) < count_old:
        j = faker.pyint(max_value=max, min_value=min)
        if must_not_contain != "" and j in not_contain_num:
            continue
        res.add(j)
    return res


def double_function(count, min, max, must_contain, must_not_contain, left_digits, right_digits):
    res = set()
    count_old = count
    # 判断约束是否正确
    if count < 0:
        return "生成个数不能为负数！"
    if min > max:
        return "最小值不能大于最大值！"
    if count > (max - min) / (1 / (pow(10, right_digits)))+1:
        return "范围内没有" + str(count_old) + "个非重复数据！"
    if max >= pow(10, left_digits):
        return "最大值与小数点左侧位数冲突"
    if left_digits + right_digits > 15:
        return "总位数不能超过15位!"

    contain_num = []
    if must_contain != "":
        contain_str = must_contain.split(',')
        contain_num = [float(i) for i in contain_str]
        for num in contain_num:
            if num > max or num < min:
                return "必须包含的数字不在约束范围内！"
        # 添加必须有的数字
        for num in contain_num:
            if num not in res:
                res.add(num)
                count = count - 1

    if must_not_contain != "":
        not_contain_str = must_not_contain.split(',')
        not_contain_num = [float(i) for i in not_contain_str]
        for num in contain_num:
            if num in not_contain_num:
                return "必须包含的数字与必须不包含的数字冲突！"

    faker = Faker(locale="zh_CN")
    # pyfloat最大值最小值参数不能是float数，需要取整
    if min >= 0:
        min_int = int(min)
        max_int = int(max + 1)
    else:
        min_int = int(min - 1)
        max_int = int(max)

    while len(res) < count_old:
        j = faker.pyfloat(max_value=max_int, min_value=min_int, left_digits=left_digits, right_digits=right_digits)
        if must_not_contain != "" and j in not_contain_num:
            continue
        if j < min or j > max:
            continue
        res.add(j)
    return res


def string_function(count, string_chars, min, max, must_contain, must_not_contain):
    res = set()
    count_old = count
    if count < 0:
        return "生成个数不能为负数！"
    if min > max:
        return "最小长度不能大于最大长度！"
    if min == max and count > 52:
        return "最多只能生成52个字符！"

    contain_str = []
    # 如果没有规定组成字符
    if string_chars == None:
        if must_contain != "":
            contain_str = must_contain.split(',')
            for str1 in contain_str:
                if len(str1) > max or len(str1) < min:
                    return "必须包含的字符串不在长度约束范围内！"
                # 添加必须有的字符串
                if str1 not in res:
                    res.add(str1)
                    count = count - 1

        if must_not_contain != "":
            not_contain_str = must_not_contain.split(',')
            for str1 in contain_str:
                if str1 in not_contain_str:
                    return "必须包含的字符串与必须不包含的字符串冲突！"

        faker = Faker(locale="zh_CN")
        while len(res) < count:
            j = faker.pystr(max_chars=max, min_chars=min)
            if must_not_contain != "" and j in not_contain_str:
                continue
            res.add(j)
        return res

    # 如果存在规定组成字符
    # 判断利用这些字符最多形成多少种字符串
    # 判断must_contain中是否有错
    if string_chars != None:
        new_chars = string_chars.split(',')
        set1 = set()
        for char in new_chars:
            set1.add(char)
        n = len(set1)
        if max != 1 and n != 1:
            if count_old > pow(n, max) - pow(n, min - 1):
                return "范围内没有这么多个非重复数据！"
        elif max != 1 and n == 1:
            if count_old > max:
                return "范围内没有这么多个非重复数据！"
        else:
            if count_old > n:
                return "范围内没有这么多个非重复数据！"

        contain_str = must_contain.split(',')
        if must_contain != "":
            for i in range(len(must_contain)):
                if must_contain[i] == ',':
                    continue
                if must_contain[i] not in new_chars:
                    return "必须存在的字符串中有非规定字符！"
            for str in contain_str:
                if len(str) > max or len(str) < min:
                    return "必须包含的字符串不在长度约束范围内！"
                # 添加必须有的字符串
                if str not in res:
                    res.add(str)
                    count = count - 1

        if must_not_contain != "":
            not_contain_str = must_not_contain.split(',')
            for str in contain_str:
                if str in not_contain_str:
                    return "必须包含的字符串与必须不包含的字符串冲突！"

        faker = Faker(locale="zh_CN")
        while len(res) < count_old:
            # 随机生成一个规定的长度
            length = faker.pyint(max_value=max, min_value=min)
            str = ""
            while len(str) < max:
                index = faker.pyint(max_value=len(new_chars) - 1, min_value=0)
                str = str + new_chars[index]
                if len(str) == length:
                    break
            if must_not_contain != "" and str in not_contain_str:
                continue
            res.add(str)
        return res


def date_function(count, start, end):
    res = set()
    if count < 0:
        return "生成个数不能为负数！"
    # 传进来的是标准形式的字符串，转化为秒级的时间戳进行随机生成
    start_ts = int(time.mktime(time.strptime(start, "%Y-%m-%d %H:%M:%S")))
    end_ts = int(time.mktime(time.strptime(end, "%Y-%m-%d %H:%M:%S")))
    if start_ts > end_ts:
        return "起始日期时间不能晚于终止日期时间！"
    if start_ts == end_ts:
        if count == 1:
            res.add(start)
            return res
        else:
            return "所选时间间隔中没有" + str(count) + "个非重复数据！"
    if count > end_ts - start_ts + 1:
        return "所选时间间隔中没有" + str(count) + "个非重复数据！"

    faker = Faker(locale="zh_CN")
    while len(res) < count:
        j_ts = faker.pyint(max_value=end_ts, min_value=start_ts)
        j = time.strftime("%Y-%m-%d %H:%M:%S", time.localtime(j_ts))
        res.add(j)
    return res


def main(type, count=10, min=0, max=10, string_chars=None, must_contain="", must_not_contain="",
         left_digits=2, right_digits=4, start="", end=""):
    if (type == "int"):
        res = int_function(count, min, max, must_contain, must_not_contain)
    elif (type == "double"):
        res = double_function(count, min, max, must_contain, must_not_contain, left_digits, right_digits)
    elif (type == "string"):
        res = string_function(count, string_chars, min, max, must_contain, must_not_contain)
    elif (type == "char"):
        res = string_function(count=count, string_chars=string_chars, min=1, max=1, must_contain=must_contain,
                              must_not_contain=must_not_contain)
    elif (type == "date"):
        res = date_function(count, start, end)
    else:
        return 0
    print(res)
    return res


def other(type, count):
    res = set()
    faker = Faker(locale="zh_CN")

    if count < 0:
        return "生成个数不能为负数！"
    if type == '国家':
        if count > 232:
            return "最多支持生成232个国家!"
        while len(res) < count:
            j = faker.country()
            res.add(j)
    elif type == '省份':
        if count > 34:
            return "最多支持生成34个省份!"
        while len(res) < count:
            j = faker.province()
            res.add(j)
    elif type == '城市':
        if count > 67:
            return "最多支持生成67个城市!"
        while len(res) < count:
            j = faker.city_name()
            res.add(j)
    elif type == '街道':
        if count > 900:
            return "最多支持生成900个街道!"
        while len(res) < count:
            j = faker.street_name()
            res.add(j)
    elif type == '姓名':
        if count > 8999:
            return "最多支持生成8999个姓名!"
        while len(res) < count:
            j = faker.name()
            res.add(j)
    elif type == '姓':
        if count > 390:
            return "最多支持生成390个姓!"
        while len(res) < count:
            j = faker.last_name()
            res.add(j)
    elif type == '名':
        if count > 130:
            return "最多支持生成130个名!"
        while len(res) < count:
            j = faker.first_name()
            res.add(j)
    elif type == '手机号':
        while len(res) < count:
            j = faker.phone_number()
            res.add(j)
    else:
        pass
    return res



class TabWidgetDemo(QTabWidget):
    def __init__(self, parent=None):
        super(TabWidgetDemo, self).__init__(parent)

        self.setWindowTitle('指定类型数据的自动生成')
        self.setWindowIcon(QIcon(":/njupt.ico"))
        self.resize(2000, 1000)
        # 限制最小的长宽
        self.setMinimumSize(1000,1000)
        # 全局变量 类型
        self.type = "int"
        self.res = set()
        # QTableView的最终父类是QWidget  将整个窗口作为一个tab

        self.font_normal = QFont()
        self.font_normal.setPointSize(13)
        # 创建多个窗口  每个窗口可以放置多个控件
        # 创建用于显示控件的窗口
        # 创建窗口tab1
        self.tab1 = QWidget()
        # 创建窗口tab2
        self.tab2 = QWidget()
        # 创建窗口tab3
        self.tab3 = QWidget()

        # 把每个窗口和选项卡绑定
        self.addTab(self.tab1, '选项卡1')
        self.addTab(self.tab2, '选项卡2')
        self.addTab(self.tab3, '选项卡3')

        # 调用
        self.tab1UI()
        self.tab2UI()
        self.tab3UI()

    # 为每个选项卡单独编写一个方法
    def tab1UI(self):
        # 创建表单布局
        layout = QFormLayout()
        # 选择类型
        type = QHBoxLayout()
        int = QRadioButton('int')
        double = QRadioButton('double')
        string = QRadioButton('string')
        char = QRadioButton('char')
        date = QRadioButton('date')
        # 初始时默认选择int
        int.setChecked(True)

        type.addWidget(int)
        type.addWidget(double)
        type.addWidget(string)
        type.addWidget(char)
        type.addWidget(date)
        type_lable = QLabel('数据类型')
        layout.addRow(type_lable, type)

        min = QLineEdit('0')
        min_lable = QLabel('最小值 (0)')
        max = QLineEdit('100')
        max_lable = QLabel('最大值 (10)')
        left_digits = QLineEdit('3')
        left_digits_lable = QLabel('小数点左侧最大位数(2)')
        right_digits = QLineEdit('4')
        right_digits_lable = QLabel('小数点右侧最大位数(4)')
        must_contain = QLineEdit('2,3')
        must_contain_lable = QLabel('必须包含的元素')
        must_not_contain = QLineEdit('4,5')
        must_not_contain_lable = QLabel('必须不包含的元素')
        string_chars = QLineEdit('a,b,c')
        string_chars_lable = QLabel('字符元素来源')
        count = QLineEdit('10')
        count_lable = QLabel('生成数据个数(10)')

        # 校验器
        intValidator = QIntValidator(self)
        countValidator = QIntValidator(self)
        countValidator.setRange(1, 99999)
        digitValidator = QIntValidator(self)
        digitValidator.setRange(1, 15)
        # todo 字符和数字要分开
        # 正则
        myValidator = QRegExpValidator(QRegExp("([-]?[0-9]+)([,，][-]?[0-9]+)*"))
        maxValidator = QRegExpValidator(QRegExp("([-]?)([0-9]+)"))
        minValidator = QRegExpValidator(QRegExp("([-]?)([0-9]+)"))

        # 绑定校验器
        left_digits.setValidator(digitValidator)
        right_digits.setValidator(digitValidator)
        count.setValidator(countValidator)
        must_contain.setValidator(myValidator)
        must_not_contain.setValidator(myValidator)
        max.setValidator(maxValidator)
        min.setValidator(minValidator)

        # 布局中添加
        layout.addRow(max_lable, max)
        layout.addRow(min_lable, min)
        # 左右位数相加，最多等于15
        layout.addRow(left_digits_lable, left_digits)
        layout.addRow(right_digits_lable, right_digits)
        layout.addRow(must_contain_lable, must_contain)
        layout.addRow(must_not_contain_lable, must_not_contain)
        layout.addRow(string_chars_lable, string_chars)
        layout.addRow(count_lable, count)

        # 初始为int，需要隐藏一些
        string_chars.hide()
        string_chars_lable.hide()
        left_digits.hide()
        left_digits_lable.hide()
        right_digits.hide()
        right_digits_lable.hide()

        # 日期相关
        start = QDateTimeEdit(QDateTime.currentDateTime())
        end = QDateTimeEdit(QDateTime.currentDateTime().addMonths(1))
        start.setDisplayFormat("yyyy-MM-dd HH:mm:ss")
        end.setDisplayFormat("yyyy-MM-dd HH:mm:ss")
        start_lable = QLabel('起始日期时间')
        end_lable = QLabel('终止日期时间')
        must_date_lable = QLabel('必须包含的时间')
        layout.addRow(start_lable, start)
        layout.addRow(end_lable, end)

        # 隐藏
        start.hide()
        start_lable.hide()
        end.hide()
        end_lable.hide()

        # 多项选择的槽函数
        int.toggled.connect(
            lambda: self.buttonState(max, min, left_digits, right_digits, string_chars, max_lable, min_lable,
                                     left_digits_lable, right_digits_lable, string_chars_lable, start, start_lable, end,
                                     end_lable, must_contain, must_contain_lable, must_not_contain,
                                     must_not_contain_lable, count))
        double.toggled.connect(
            lambda: self.buttonState(max, min, left_digits, right_digits, string_chars, max_lable, min_lable,
                                     left_digits_lable, right_digits_lable, string_chars_lable, start, start_lable, end,
                                     end_lable, must_contain, must_contain_lable, must_not_contain,
                                     must_not_contain_lable, count))
        string.toggled.connect(
            lambda: self.buttonState(max, min, left_digits, right_digits, string_chars, max_lable, min_lable,
                                     left_digits_lable, right_digits_lable, string_chars_lable, start, start_lable, end,
                                     end_lable, must_contain, must_contain_lable, must_not_contain,
                                     must_not_contain_lable, count))
        char.toggled.connect(
            lambda: self.buttonState(max, min, left_digits, right_digits, string_chars, max_lable, min_lable,
                                     left_digits_lable, right_digits_lable, string_chars_lable, start, start_lable, end,
                                     end_lable, must_contain, must_contain_lable, must_not_contain,
                                     must_not_contain_lable, count))
        date.toggled.connect(
            lambda: self.buttonState(max, min, left_digits, right_digits, string_chars, max_lable, min_lable,
                                     left_digits_lable, right_digits_lable, string_chars_lable, start, start_lable, end,
                                     end_lable, must_contain, must_contain_lable, must_not_contain,
                                     must_not_contain_lable, count))

        button = QPushButton('生成数据')
        layout.addRow(button)

        # 导出的按钮
        out = QHBoxLayout()
        out_txt = QPushButton('导出txt')
        out_excel = QPushButton('导出Excel')
        out_word = QPushButton('导出Word')
        out.addWidget(out_txt)
        out.addWidget(out_excel)
        out.addWidget(out_word)
        layout.addRow(out)

        txt = QTextEdit()
        layout.addRow(txt)

        button.clicked.connect(
            lambda: self.onClick_Button(min, max, left_digits, right_digits, must_contain, must_not_contain,
                                        string_chars, start, end, count, txt))
        out_excel.clicked.connect(lambda: self.onClick_Excel(count))
        out_txt.clicked.connect(lambda: self.onClick_txt(count))
        out_word.clicked.connect(lambda: self.onClick_word(count))

        self.setTabText(0, '基本类型')
        self.setFont(self.font_normal)

        # 字体大小的设置
        txt.setFontPointSize(14)

        # 美化
        button.setMinimumSize(50,70)
        out_txt.setMinimumSize(50,70)
        out_excel.setMinimumSize(50,70)
        out_word.setMinimumSize(50,70)

        type_lable.setFont(self.font_normal)
        int.setFont(self.font_normal)
        double.setFont(self.font_normal)
        string.setFont(self.font_normal)
        char.setFont(self.font_normal)
        date.setFont(self.font_normal)
        min.setFont(self.font_normal)
        min.setFont(self.font_normal)
        min_lable.setFont(self.font_normal)
        max.setFont(self.font_normal)
        max_lable.setFont(self.font_normal)
        left_digits.setFont(self.font_normal)
        left_digits_lable.setFont(self.font_normal)
        right_digits.setFont(self.font_normal)
        right_digits_lable.setFont(self.font_normal)
        must_contain.setFont(self.font_normal)
        must_not_contain.setFont(self.font_normal)
        must_contain_lable.setFont(self.font_normal)
        must_not_contain_lable.setFont(self.font_normal)
        count.setFont(self.font_normal)
        count_lable.setFont(self.font_normal)

        button.setFont(self.font_normal)
        out_txt.setFont(self.font_normal)
        out_excel.setFont(self.font_normal)

        # todo tab按钮切换时执行方法
        self.currentChanged.connect(lambda: self.tabchanged(int, double, string, char, date))

        # 装载
        self.tab1.setLayout(layout)

    def tab2UI(self):
        layout = QFormLayout()
        type = QHBoxLayout()
        country = QRadioButton('国家')
        province = QRadioButton('省份')
        city = QRadioButton('城市')
        street = QRadioButton('街道')
        type.addWidget(country)
        type.addWidget(province)
        type.addWidget(city)
        type.addWidget(street)
        layout.addRow(QLabel('类别'), type)

        country.setChecked(True)

        count = QLineEdit()
        intvalidator = QIntValidator()
        intvalidator.setRange(1, 1000)
        count.setValidator(intvalidator)

        layout.addRow('生成数据个数(10)', count)

        button = QPushButton('生成数据')
        layout.addRow(button)

        # 导出的按钮
        out = QHBoxLayout()
        out_txt = QPushButton('导出txt')
        out_excel = QPushButton('导出Excel')
        out_word = QPushButton('导出Word')
        out.addWidget(out_txt)
        out.addWidget(out_excel)
        out.addWidget(out_word)
        layout.addRow(out)

        txt = QTextEdit()
        layout.addRow(txt)
        # 字体大小的设置
        txt.setFontPointSize(14)

        # 按钮与type的连接
        country.toggled.connect(self.buttonStateOther)
        province.toggled.connect(self.buttonStateOther)
        city.toggled.connect(self.buttonStateOther)
        street.toggled.connect(self.buttonStateOther)

        button.clicked.connect(
            lambda: self.onClick_Button_other(count, txt))
        out_excel.clicked.connect(lambda: self.onClick_Excel(count))
        out_txt.clicked.connect(lambda: self.onClick_txt(count))
        out_word.clicked.connect(lambda: self.onClick_word(count))

        button.setMinimumSize(50, 70)
        out_txt.setMinimumSize(50, 70)
        out_excel.setMinimumSize(50, 70)
        out_word.setMinimumSize(50, 70)

        # todo tab按钮切换时执行方法
        self.currentChanged.connect(lambda: self.tabchanged(country, province, city, street, None))

        self.setTabText(1, "地理相关")
        self.tab2.setLayout(layout)

    def tab3UI(self):
        layout = QFormLayout()
        type = QHBoxLayout()
        country = QRadioButton('姓名')
        province = QRadioButton('姓')
        city = QRadioButton('名')
        street = QRadioButton('手机号')
        type.addWidget(country)
        type.addWidget(province)
        type.addWidget(city)
        type.addWidget(street)
        layout.addRow(QLabel('类别'), type)

        country.setChecked(True)

        count = QLineEdit()
        intvalidator = QIntValidator()
        intvalidator.setRange(1, 1000)
        count.setValidator(intvalidator)

        layout.addRow('生成数据个数(10)', count)

        button = QPushButton('生成数据')
        layout.addRow(button)

        # 导出的按钮
        out = QHBoxLayout()
        out_txt = QPushButton('导出txt')
        out_excel = QPushButton('导出Excel')
        out_word = QPushButton('导出Word')
        out.addWidget(out_txt)
        out.addWidget(out_excel)
        out.addWidget(out_word)
        layout.addRow(out)

        txt = QTextEdit()
        layout.addRow(txt)
        # 字体大小的设置
        txt.setFontPointSize(14)

        # 按钮与type的连接
        country.toggled.connect(self.buttonStateOther)
        province.toggled.connect(self.buttonStateOther)
        city.toggled.connect(self.buttonStateOther)
        street.toggled.connect(self.buttonStateOther)

        button.clicked.connect(
            lambda: self.onClick_Button_other(count, txt))
        out_excel.clicked.connect(lambda: self.onClick_Excel(count))
        out_txt.clicked.connect(lambda: self.onClick_txt(count))
        out_word.clicked.connect(lambda: self.onClick_word(count))

        button.setMinimumSize(50, 70)
        out_txt.setMinimumSize(50, 70)
        out_excel.setMinimumSize(50, 70)
        out_word.setMinimumSize(50, 70)

        # todo tab按钮切换时执行方法
        self.currentChanged.connect(lambda: self.tabchanged(country, province, city, street, None))

        self.setTabText(2, "人物相关")
        self.tab3.setLayout(layout)

    # tab页切换时设定对应的type
    def tabchanged(self, but_1, but_2, but_3, but_4, but_5):
        if self.currentIndex() == 0:
            if but_1.isChecked() and but_1.text() == 'int':
                self.type = 'int'
            elif but_2.isChecked() and but_2.text() == 'double':
                self.type = 'double'
            elif but_3.isChecked() and but_3.text() == 'string':
                self.type = 'string'
            elif but_4.isChecked() and but_4.text() == 'char':
                self.type = 'char'
            elif but_5 != None and but_5.isChecked() and but_5.text() == 'date':
                self.type = 'date'
            else:
                print()
        elif self.currentIndex() == 1:
            if but_1.isChecked() and but_1.text() == '国家':
                self.type = '国家'
            elif but_2.isChecked() and but_2.text() == '省份':
                self.type = '省份'
            elif but_3.isChecked() and but_3.text() == '城市':
                self.type = '城市'
            elif but_4.isChecked() and but_4.text() == '街道':
                self.type = '街道'
            else:
                print()
        else:
            if but_1.isChecked() and but_1.text() == '姓名':
                self.type = '姓名'
            elif but_2.isChecked() and but_2.text() == '姓':
                self.type = '姓'
            elif but_3.isChecked() and but_3.text() == '名':
                self.type = '名'
            elif but_4.isChecked() and but_4.text() == '手机号':
                self.type = '手机号'
            else:
                print()

    def onClick_Button(self, Qmin, Qmax, Qleft_digits, Qright_digits, Qmust_contain, Qmust_not_contain, Qstring_chars,
                       Qstart, Qend, Qcount, txt):
        if Qmin.text() == '':
            min = 0
        else:
            if self.type == "int":
                min = int(Qmin.text())
            else:
                min = float(Qmin.text())
        if Qmax.text() == '':
            max = 10
        else:
            if self.type == "int":
                max = int(Qmax.text())
            else:
                max = float(Qmax.text())
        if Qcount.text() == '':
            count = 10
        else:
            count = int(Qcount.text())
        if Qleft_digits.text() == '':
            left_digits = 2
        else:
            left_digits = int(Qleft_digits.text())
        if Qright_digits.text() == '':
            right_digits = 4
        else:
            right_digits = int(Qright_digits.text())

        # 如果有中文逗号,就替换
        table = {ord(f): ord(t) for f, t in zip(
            u'，。！？【】（）％＃＠＆１２３４５６７８９０',
            u',.!?[]()%#@&1234567890')}
        must_contain = Qmust_contain.text()
        must_not_contain = Qmust_not_contain.text()
        must_contain = must_contain.translate(table)
        must_not_contain = must_not_contain.translate(table)


        if Qstring_chars.text() == '':
            string_chars = None
        else:
            string_chars = Qstring_chars.text().translate(table)

        start = Qstart.text()
        end = Qend.text()
        self.res = main(self.type, count, min, max, string_chars, must_contain, must_not_contain, left_digits,
                             right_digits, start, end)
        txt.setText("")
        txt.setText(str(self.res).strip('{').strip('}'))

    def onClick_Button_other(self, Qcount, txt):
        if Qcount.text() == '':
            count = 10
        else:
            count = int(Qcount.text())

        self.res = other(self.type, count)
        txt.setText("")
        txt.setText(str(self.res).strip('{').strip('}'))

    def onClick_Excel(self, Qcount):
        if len(self.res) == 0:
            # 消息框提示
            msg_box = QMessageBox(QMessageBox.Information, '提示', '请先生成数据！')
            msg_box.exec_()
            return None

        wb = Workbook()
        sheet = wb.active
        ws = wb['Sheet']
        # 设置列宽
        sheet.column_dimensions['A'].width = 35
        for i in iter(self.res):
            sheet.append([i])

        if Qcount.text() == '':
            count = 10
        else:
            count = int(Qcount.text())

        # 单元格居中
        align = Alignment(horizontal='center', vertical='center')
        # openpyxl的下标从1开始
        for i in range(1, count + 1):
            ws.cell(i, 1).alignment = align
        # 获取当前时间，用于命名
        time_cur = time.strftime('%Y-%m-%d-%H-%M-%S', time.localtime(time.time()))
        wb.save('out-' + self.type + '-' + time_cur + '.xlsx')
        # 消息框提示
        msg_box = QMessageBox(QMessageBox.Information, '提示', '导出Excel成功！')
        msg_box.exec_()

    def onClick_txt(self, Qcount):
        if len(self.res) == 0:
            # 消息框提示
            msg_box = QMessageBox(QMessageBox.Information, '提示', '请先生成数据！')
            msg_box.exec_()
            return None

        # 获取当前时间，用于命名
        time_cur = time.strftime('%Y-%m-%d-%H-%M-%S', time.localtime(time.time()))
        f = open('out-' + self.type + '-' + time_cur + '.txt', 'w')
        for i in iter(self.res):
            f.write(str(i))
            f.write('\n')
        f.close()
        # 消息框提示
        msg_box = QMessageBox(QMessageBox.Information, '提示', '导出txt成功！')
        msg_box.exec_()

    def onClick_word(self, Qcount):
        if len(self.res) == 0:
            # 消息框提示
            msg_box = QMessageBox(QMessageBox.Information, '提示', '请先生成数据！')
            msg_box.exec_()
            return None
        if Qcount.text() == '':
            count = 10
        else:
            count = int(Qcount.text())

        if count > 15000:
            # 消息框提示
            msg_box = QMessageBox(QMessageBox.Information, '提示', "Word最大只支持15000条数据！")
            msg_box.exec_()
            return

        # 获取当前时间，用于命名
        time_cur = time.strftime('%Y-%m-%d-%H-%M-%S', time.localtime(time.time()))
        # 打开一个word文档
        document = Document()
        for i in iter(self.res):
            document.add_paragraph(str(i))
        document.save('out-' + self.type + '-' + time_cur + '.docx')
        # 消息框提示
        msg_box = QMessageBox(QMessageBox.Information, '提示', '导出Word成功！')
        msg_box.exec_()

    def buttonState(self, max, min, left_digits, right_digits, string_chars, max_lable, min_lable,
                    left_digits_lable, right_digits_lable, string_chars_lable, start, start_lable, end,
                    end_lable, must_contain, must_contain_lable, must_not_contain, must_not_contain_lable, count):
        # sander为事件发生者
        if self.sender().isChecked() == True:
            self.type = self.sender().text()

        if self.type == 'int':
            left_digits.hide()
            left_digits_lable.hide()
            right_digits.hide()
            right_digits_lable.hide()
            string_chars.hide()
            string_chars_lable.hide()
            max.show()
            max_lable.show()
            min.show()
            min_lable.show()
            start.hide()
            start_lable.hide()
            end.hide()
            end_lable.hide()
            must_contain.show()
            must_contain_lable.show()
            must_not_contain.show()
            must_not_contain_lable.show()

            # 清空
            min.setText('')
            max.setText('')
            left_digits.setText('')
            right_digits.setText('')
            must_contain.setText('')
            must_not_contain.setText('')
            string_chars.setText('')
            # 正则表达式限制
            myValidator = QRegExpValidator(QRegExp("([-]?[0-9]+)([,，][-]?[0-9]+)*"))
            must_contain.setValidator(myValidator)
            must_not_contain.setValidator(myValidator)
            maxValidator = QRegExpValidator(QRegExp("([-]?)([0-9]+)"))
            minValidator = QRegExpValidator(QRegExp("([-]?)([0-9]+)"))
            max.setValidator(maxValidator)
            min.setValidator(minValidator)


        elif self.type == 'double':
            left_digits.show()
            left_digits_lable.show()
            right_digits.show()
            right_digits_lable.show()
            string_chars.hide()
            string_chars_lable.hide()
            max.show()
            max_lable.show()
            min.show()
            min_lable.show()
            start.hide()
            start_lable.hide()
            end.hide()
            end_lable.hide()
            must_contain.show()
            must_contain_lable.show()
            must_not_contain.show()
            must_not_contain_lable.show()

            # 清空
            min.setText('')
            max.setText('')
            left_digits.setText('')
            right_digits.setText('')
            must_contain.setText('')
            must_not_contain.setText('')
            string_chars.setText('')
            # 正则表达式限制
            myValidator = QRegExpValidator(QRegExp("([-]?[0-9]+[.]?[0-9]+)([,，]([-])?[0-9]+[.]?[0-9]+)*"))
            must_contain.setValidator(myValidator)
            must_not_contain.setValidator(myValidator)
            maxValidator = QRegExpValidator(QRegExp("[-]?([0-9]+)[.]([0-9]+)"))
            minValidator = QRegExpValidator(QRegExp("[-]?([0-9]+)[.]([0-9]+)"))
            max.setValidator(maxValidator)
            min.setValidator(minValidator)

        elif self.type == 'string':
            left_digits.hide()
            left_digits_lable.hide()
            right_digits.hide()
            right_digits_lable.hide()
            string_chars.show()
            string_chars_lable.show()
            max.show()
            max_lable.show()
            min.show()
            min_lable.show()
            start.hide()
            start_lable.hide()
            end.hide()
            end_lable.hide()
            must_contain.show()
            must_contain_lable.show()
            must_not_contain.show()
            must_not_contain_lable.show()

            # 清空
            min.setText('1')
            max.setText('')
            left_digits.setText('')
            right_digits.setText('')
            must_contain.setText('')
            must_not_contain.setText('')
            string_chars.setText('')
            # 正则表达式限制
            myValidator = QRegExpValidator(QRegExp("([a-zA-Z0-9]+)([,，][a-zA-Z0-9]+)*"))
            must_contain.setValidator(myValidator)
            must_not_contain.setValidator(myValidator)
            charValidator = QRegExpValidator(QRegExp("([a-zA-Z0-9])([,，][a-zA-Z0-9])*"))
            string_chars.setValidator(charValidator)
            maxValidator = QRegExpValidator(QRegExp("([0-9]+)"))
            minValidator = QIntValidator(self)
            minValidator.setRange(1,99999)
            max.setValidator(maxValidator)
            min.setValidator(minValidator)

        elif self.type == 'char':
            left_digits.hide()
            left_digits_lable.hide()
            right_digits.hide()
            right_digits_lable.hide()
            string_chars.show()
            string_chars_lable.show()
            min.hide()
            min_lable.hide()
            max.hide()
            max_lable.hide()
            start.hide()
            start_lable.hide()
            end.hide()
            end_lable.hide()
            must_contain.show()
            must_contain_lable.show()
            must_not_contain.show()
            must_not_contain_lable.show()

            # 清空
            min.setText('')
            max.setText('')
            left_digits.setText('')
            right_digits.setText('')
            must_contain.setText('')
            must_not_contain.setText('')
            string_chars.setText('')
            # 正则表达式限制
            myValidator = QRegExpValidator(QRegExp("([a-zA-Z0-9])([,，][a-zA-Z0-9])*"))
            must_contain.setValidator(myValidator)
            must_not_contain.setValidator(myValidator)
            charValidator = QRegExpValidator(QRegExp("([a-zA-Z0-9])([,，][a-zA-Z0-9])*"))
            string_chars.setValidator(charValidator)

        else:  # date
            start.show()
            start_lable.show()
            end.show()
            end_lable.show()
            left_digits.hide()
            left_digits_lable.hide()
            right_digits.hide()
            right_digits_lable.hide()
            string_chars.hide()
            string_chars_lable.hide()
            max.hide()
            max_lable.hide()
            min.hide()
            min_lable.hide()
            must_contain.hide()
            must_contain_lable.hide()
            must_not_contain.hide()
            must_not_contain_lable.hide()

            # 清空
            min.setText('')
            max.setText('')
            left_digits.setText('')
            right_digits.setText('')
            must_contain.setText('')
            must_not_contain.setText('')
            string_chars.setText('')

    def buttonStateOther(self):
        # sander为事件发生者
        if self.sender().isChecked() == True:
            self.type = self.sender().text()


def read_qss_file(qss_file_name):
    with open(qss_file_name, 'r',  encoding='UTF-8') as file:
        return file.read()


if __name__ == '__main__':
    app = QApplication(sys.argv)
    demo = TabWidgetDemo()
    # # 设置qss美化
    # style_file = './qss/flatgray.qss'
    # style_sheet = read_qss_file(style_file)
    # demo.setStyleSheet(style_sheet)
    # extra = {
    #     'font_size': '14px'
    # }
    # apply_stylesheet(app, theme='light_teal.xml', extra=extra)

    demo.show()
    sys.exit(app.exec_())